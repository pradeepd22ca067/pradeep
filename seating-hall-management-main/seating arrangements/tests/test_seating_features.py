import os
import sqlite3
import tempfile
import unittest

from openpyxl import load_workbook

import zen_app


class SeatingFeatureTests(unittest.TestCase):
    def setUp(self):
        self.database = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.database.close()
        self.original_database_path = zen_app.app.config.get("DATABASE_PATH")
        zen_app.app.config.update(TESTING=True, DATABASE_PATH=self.database.name)
        connection = sqlite3.connect(self.database.name)
        connection.execute(
            "CREATE TABLE room (room_no INTEGER, col INTEGER, row INTEGER, seat INTEGER, session TEXT)"
        )
        connection.execute("INSERT INTO room VALUES (777, 2, 2, 4, NULL)")
        connection.commit()
        connection.close()
        zen_app.initialize_history_table()
        self.client = zen_app.app.test_client()
        self.generated_file = os.path.join(
            zen_app.APP_DIR, "static", "excel",
            "seating_777_FN_2099-12-31_I.xlsx"
        )

    def tearDown(self):
        zen_app.app.config["DATABASE_PATH"] = self.original_database_path
        if os.path.exists(self.generated_file):
            os.remove(self.generated_file)
        os.remove(self.database.name)

    def generation_form(self, **overrides):
        form = {
            "room": "777",
            "date": "2099-12-31",
            "session": "FN",
            "shift": "I",
            "bca_start": "101",
            "bca_end": "103",
            "subject_code_bca": "BCA101",
            "bsc_start": "",
            "bsc_end": "",
            "subject_code_bsc": "BSC101",
            "bcom_start": "",
            "bcom_end": "",
            "subject_code_bcom": "BCOM101",
            "bba_start": "",
            "bba_end": "",
            "subject_code_bba": "BBA101",
            "missing": "",
        }
        form.update(overrides)
        return form

    def test_rejects_student_count_over_room_capacity(self):
        response = self.client.post(
            "/Generate",
            data=self.generation_form(bca_start="101", bca_end="105"),
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn(b"5 students exceed this room", response.data)

    def test_rejects_overlapping_roll_number_ranges(self):
        response = self.client.post(
            "/Generate",
            data=self.generation_form(
                bca_start="101", bca_end="102",
                bsc_start="102", bsc_end="103",
            ),
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn(b"ranges overlap", response.data)

    def test_saves_assignment_for_history_lookup_and_print(self):
        response = self.client.post("/Generate", data=self.generation_form())
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.content_type,
        )
        response.close()
        workbook = load_workbook(self.generated_file, read_only=True)
        self.assertIn("Seating Plan", workbook.sheetnames)
        workbook.close()

        history = self.client.get("/history")
        self.assertEqual(history.status_code, 200)
        self.assertIn(b"2099-12-31", history.data)

        lookup = self.client.post("/lookup", data={"roll_no": "101"})
        self.assertEqual(lookup.status_code, 200)
        self.assertIn(b"Room 777", lookup.data)
        self.assertIn(b"Row", lookup.data)

        arrangement_id = self.client.get("/history").data.decode().split(
            "/history/", 1
        )[1].split('"', 1)[0]
        printable = self.client.get(f"/history/{arrangement_id}/print")
        self.assertEqual(printable.status_code, 200)
        self.assertIn(b"Save as PDF", printable.data)
        self.assertIn(b"101", printable.data)


if __name__ == "__main__":
    unittest.main()
